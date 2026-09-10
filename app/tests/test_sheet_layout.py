"""The shared sheet layout: the on-screen worksheet and the downloaded workbook
are the same rows.

`sheet_layout` owns the row order, the keys and the arithmetic; the workbook
renders it. The first test here is the anti-drift contract: every defined name
the workbook writes sits at the (column, row) the layout says, so a grid that
numbers its rows from the layout is numbering them the way the .xlsx does.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services import business_statement_schema as bss
from app.services import financial_templates_xlsx as templates
from app.services import pfs_schema
from app.services import sheet_layout as sl
from app.services.bucket_ai import MAX_SPREADSHEET_ROWS
from app.services.financial_statements import DEBT_COLUMN_LABELS, DEBT_COLUMNS

KINDS = sl.KINDS


def _workbook(kind):
    return load_workbook(BytesIO(templates.build_workbook(kind)))


def _empty_body(kind):
    if kind in bss.KINDS:
        return bss.empty_body(kind)
    if kind == "pfs":
        return pfs_schema.empty_body()
    return {"debts": []}


# --- The contract ------------------------------------------------------------


@pytest.mark.parametrize("kind", KINDS)
def test_the_layout_and_the_workbook_agree_row_for_row(kind):
    """For every defined name in the built workbook, (column, row) is the
    layout cell's (c, r), matched on xlsx_name — and the other way round."""
    wb = _workbook(kind)
    sheet = sl.layout(kind)
    prefix = sheet.name_prefix + "."
    assert wb.defined_names
    for name, defined in wb.defined_names.items():
        assert name.startswith(prefix), name
        row, cell = sheet.by_xlsx_name[name[len(prefix):]]
        address = defined.attr_text.rsplit("!", 1)[1].replace("$", "")
        assert address == f"{get_column_letter(cell.c)}{row.r}", name
    for suffix in sheet.by_xlsx_name:
        assert prefix + suffix in wb.defined_names, suffix

    ws = wb.active
    for row in sheet.rows:
        for cell in row.cells:
            written = ws.cell(row=row.r, column=cell.c).value
            if cell.type == "label":
                assert written == cell.label, (row.r, cell.c)
            elif cell.type == "formula":
                assert isinstance(written, str) and written.startswith("="), (row.r, cell.c)
            else:
                assert written is None, (row.r, cell.c)
    assert [row.r for row in sheet.rows] == sorted({row.r for row in sheet.rows})
    assert sheet.rows[-1].r == ws.max_row


@pytest.mark.parametrize("kind", KINDS)
def test_every_schema_key_is_exactly_one_layout_cell(kind):
    sheet = sl.layout(kind)
    names = [cell.xlsx_name for row in sheet.rows for cell in row.cells if cell.xlsx_name]
    keys = [cell.key for row in sheet.rows for cell in row.cells if cell.key]
    assert len(set(names)) == len(names)
    assert len(set(keys)) == len(keys)
    for key in templates.schema_keys(kind):
        assert names.count(key) == 1, key


@pytest.mark.parametrize("kind", KINDS)
def test_every_compute_key_the_layout_declares_is_a_totals_entry(kind):
    computed = sl.totals(kind, _empty_body(kind))
    keys = sl.compute_keys(kind)
    assert keys
    for key in keys:
        assert key in computed, key
    sheet = sl.layout(kind)
    for row in sheet.rows:
        for cell in row.cells:
            if cell.type == "formula":
                assert (cell.compute is None) != (cell.source is None), cell.xlsx_name
                if cell.source is not None:
                    assert cell.source in sheet.by_key, cell.source


@pytest.mark.parametrize("kind", KINDS)
def test_formulas_reference_only_cells_the_sheet_has(kind):
    sheet = sl.layout(kind)
    known = set(sheet.by_key) | set(sheet.by_xlsx_name)
    for row in sheet.rows:
        for cell in row.cells:
            if cell.formula:
                for ref in templates._FORMULA_REF.findall(cell.formula):
                    assert ref in known, (cell.xlsx_name, ref)


# --- Reading and writing bodies ----------------------------------------------


def _filled_values(sheet):
    return {cell.key: str(100 + index) for index, (_, cell) in enumerate(sl.input_cells(sheet))}


@pytest.mark.parametrize("kind", KINDS)
def test_flatten_unflatten_round_trips_a_filled_body(kind):
    values = _filled_values(sl.layout(kind))
    body = sl.unflatten(kind, values, _empty_body(kind))
    assert sl.flatten(kind, body) == values
    # And a second pass over the filled body changes nothing.
    assert sl.flatten(kind, sl.unflatten(kind, sl.flatten(kind, body), body)) == values


def test_unflatten_lands_where_the_schemas_read():
    pl = sl.unflatten("p_and_l", {"gross_revenue": "1,000", "cost_of_goods_sold": "250", "business_name": "Acme"}, bss.empty_body("p_and_l"))
    assert pl["sections"]["revenue"]["gross_revenue"] == "1,000"
    assert pl["header"]["business_name"] == "Acme"
    assert bss.totals("p_and_l", pl)["gross_profit"] == 750

    pfs = sl.unflatten(
        "pfs",
        {"cash_on_hand": "10", "name": "Ann", "notes_payable.r1.current_balance": "5", "statement_date": "2026-01-31"},
        pfs_schema.empty_body(),
    )
    assert pfs["assets"]["cash_on_hand"] == "10" and pfs["applicant"]["name"] == "Ann"
    assert pfs["schedules"]["notes_payable"] == [{"current_balance": "5"}]   # keyed, never by label
    assert pfs[sl.PFS_AS_OF_KEY] == "2026-01-31"
    assert pfs_schema.totals(pfs)["total_assets"] == 10

    ds = sl.unflatten("debt_schedule", {"r2.lender": "Bank", "business_name": "Acme"}, {"debts": []})
    assert ds["debts"] == [{}, {"lender": "Bank"}] and ds["business_name"] == "Acme"


def test_blank_stays_null_and_unknown_keys_are_refused():
    base = bss.empty_body("p_and_l")
    base["sections"]["revenue"]["gross_revenue"] = "5"
    body = sl.unflatten("p_and_l", {"gross_revenue": ""}, base)
    assert body["sections"]["revenue"]["gross_revenue"] is None
    assert base["sections"]["revenue"]["gross_revenue"] == "5"   # a copy, not the base
    assert sl.flatten("p_and_l", body)["gross_revenue"] == ""
    with pytest.raises(KeyError):
        sl.unflatten("p_and_l", {"gross_profit": "1"}, base)   # a formula, not an input
    with pytest.raises(KeyError):
        sl.unflatten("debt_schedule", {"r1.nope": "x"}, {"debts": []})
    with pytest.raises(KeyError):
        sl.unflatten("pfs", {"real_estate.r1.nope": "x"}, pfs_schema.empty_body())


def test_list_rows_are_keyed_by_identity_when_the_body_has_ids():
    body = {
        "debts": [
            {"id": "a1", "lender": "First", "balance": "100", "monthly_payment": "10"},
            {"id": "b2", "lender": "Second", "balance": "1,050", "monthly_payment": ""},
        ]
    }
    flat = sl.flatten("debt_schedule", body)
    assert flat["a1.lender"] == "First" and flat["b2.balance"] == "1,050" and flat["b2.monthly_payment"] == ""
    assert "r1.lender" not in flat
    rebuilt = sl.unflatten("debt_schedule", flat, {"debts": []})
    assert [row["id"] for row in rebuilt["debts"]] == ["a1", "b2"]
    assert sl.flatten("debt_schedule", rebuilt) == flat
    # An edit to one cell of a known row touches nothing else.
    patched = sl.unflatten("debt_schedule", {"b2.lender": "Renamed"}, body)
    assert patched["debts"][1] == {"id": "b2", "lender": "Renamed", "balance": "1,050", "monthly_payment": ""}
    assert patched["debts"][0] == body["debts"][0]

    layout = sl.layout_for_body("debt_schedule", body)
    data = [row for row in layout.rows if row.kind == "data"]
    assert [row.row_key for row in data] == ["a1", "b2"]
    assert data[1].cells[0].key == "b2.lender" and data[1].cells[0].xlsx_name == "r2.lender"
    assert data[0].r == 7   # the row the workbook writes, hidden key row and all
    total = next(cell for row in layout.rows if row.kind == "formula" for cell in row.cells if cell.compute == "total_balance")
    assert total.formula == "=SUM({a1.balance}:{b2.balance})"
    assert sl.debt_totals(body) == {"total_balance": 1150, "total_monthly_payment": 10}


def test_pfs_schedule_rows_carry_ids_the_same_way_and_legacy_rows_are_read_by_label():
    body = pfs_schema.empty_body()
    body["schedules"]["real_estate"] = [
        # A row stored before the column keys existed: keyed by label.
        {"id": "h1", "Property address": "1 Main St", "Mortgage balance": "90"},
        {"id": "h2", "property_address": "2 Side St"},
    ]
    flat = sl.flatten("pfs", body)
    assert flat["real_estate.h1.property_address"] == "1 Main St"
    assert flat["real_estate.h1.mortgage_balance"] == "90"
    assert flat["real_estate.h2.property_address"] == "2 Side St"
    assert "notes_payable.r1.current_balance" not in flat   # an empty schedule has no rows
    rebuilt = sl.unflatten("pfs", flat, pfs_schema.empty_body())
    assert sl.flatten("pfs", rebuilt) == flat
    assert set(rebuilt["schedules"]["real_estate"][0]) == {"id", *pfs_schema.SCHEDULES_BY_KEY["real_estate"].column_keys}
    assert body["schedules"]["real_estate"][0]["Property address"] == "1 Main St"   # nothing written back
    layout = sl.layout_for_body("pfs", body)
    data = [row for row in layout.rows if row.kind == "data"]
    assert [row.block for row in data] == ["real_estate", "real_estate"]
    assert data[0].cells[0].xlsx_name == "real_estate.r1.property_address"
    assert data[1].cells[0].key == "real_estate.h2.property_address"


# --- The PFS schedules -------------------------------------------------------


def test_the_pfs_layout_carries_all_eight_schedules():
    sheet = sl.layout("pfs")
    blocks = [row.block for row in sheet.rows if row.kind == "heading" and row.block]
    assert blocks == [spec.key for spec in pfs_schema.SCHEDULES]
    assert len(blocks) == 8
    for spec in pfs_schema.SCHEDULES:
        colhead = next(row for row in sheet.rows if row.kind == "colhead" and row.block == spec.key)
        assert [cell.label for cell in colhead.cells] == [label for _, _, label in sl.schedule_columns(spec)]
        data = [row for row in sheet.rows if row.kind == "data" and row.block == spec.key]
        assert len(data) == sl.PFS_SCHEDULE_ROWS
        for row in data:
            assert len(row.cells) == len(spec.columns)
            assert all(cell.c in sl.PFS_SCHEDULE_COLUMNS for cell in row.cells)
            assert all(cell.c != 3 for cell in row.cells)   # the hidden key column
            assert all(cell.editable and cell.key and cell.xlsx_name for cell in row.cells)
    assert sheet.rows[-1].r <= MAX_SPREADSHEET_ROWS
    assert sheet.rows[-1].kind == "note"


def test_schedule_columns_read_the_schema_keys_and_fall_back_to_label_slugs():
    spec = pfs_schema.SCHEDULES_BY_KEY["other_liabilities"]
    assert sl.schedule_columns(spec) == [("description", "description", "Description"), ("amount", "amount", "Amount")]
    # The key the schema chose is the slug of its label, so a name never moves.
    for spec in pfs_schema.SCHEDULES:
        for key, label in spec.fields:
            assert sl._slug(label) == key, (spec.key, label)

    class _Spec:
        key = "x"
        columns = ("Name and address of noteholder", "Amount")

    assert sl.schedule_columns(_Spec()) == [
        ("Name and address of noteholder", "name_and_address_of_noteholder", "Name and address of noteholder"),
        ("Amount", "amount", "Amount"),
    ]


# --- The debt schedule -------------------------------------------------------


def test_the_debt_schedule_layout_is_the_template_grid():
    sheet = sl.layout("debt_schedule")
    assert [column.label for column in sheet.columns] == list(DEBT_COLUMN_LABELS)
    colhead = next(row for row in sheet.rows if row.kind == "colhead")
    assert colhead.r == 5
    data = [row for row in sheet.rows if row.kind == "data"]
    assert len(data) == sl.DEBT_ROWS and data[0].r == 7 and data[-1].r == 21
    assert [cell.key for cell in data[6].cells] == [f"r7.{column}" for column in DEBT_COLUMNS]
    secured = next(cell for cell in data[0].cells if cell.key == "r1.secured")
    assert secured.type == "select" and secured.options == ("secured", "unsecured")
    assert sheet.freeze == (5, 1)
    assert sl.compute_keys("debt_schedule") == ["total_balance", "total_monthly_payment"]


def test_layouts_are_versioned_and_titled_like_the_workbook():
    for kind in KINDS:
        sheet = sl.layout(kind)
        assert sheet.layout_version == sl.LAYOUT_VERSION
        assert sheet.title == templates.SHEET_TITLES[kind]
        assert sheet.name_prefix == templates.NAME_PREFIX[kind]
        assert sheet.schema_version == sl.SCHEMA_VERSIONS[kind]
    with pytest.raises(KeyError):
        sl.layout("tax_return")
