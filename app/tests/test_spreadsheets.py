"""One reader for .xlsx, legacy .xls and .csv: the AI and the review screens
see the same rows, and what was left out is said plainly.
"""

from __future__ import annotations

import re
import zipfile
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services import spreadsheets
from app.services.bucket_ai import (
    _append_review_file_content,
    _extract_xls_text,
    _spreadsheet_skip_still_applies,
    detect_document_content_type,
)


def _xlsx(rows: list[list[object]], *, title: str = "Sheet1") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    raw = BytesIO()
    workbook.save(raw)
    return raw.getvalue()


def _without_dimension(xlsx: bytes) -> bytes:
    """A workbook the way web exports and Numbers save it: no <dimension> record."""

    source = zipfile.ZipFile(BytesIO(xlsx))
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(rb"<dimension ref=\"[^\"]*\"\s*/>", b"", data)
                assert b"<dimension" not in data
            target.writestr(item, data)
    return output.getvalue()


# --- .xlsx -----------------------------------------------------------------------


def test_a_workbook_saved_without_a_dimension_record_is_still_read() -> None:
    raw = _without_dimension(_xlsx([["Account", "Amount"], ["Gross sales", 1250000.0], ["COGS", 850000]]))

    workbook = spreadsheets.read_workbook(raw, "xlsx")

    assert workbook.kind == "xlsx"
    assert workbook.sheets[0].rows == [["Account", "Amount"], ["Gross sales", "1250000"], ["COGS", "850000"]]
    assert workbook.sheets[0].total_rows == 3
    assert workbook.sheets[0].warnings == []


def test_rows_beyond_the_limit_are_counted_and_reported_not_silently_dropped() -> None:
    raw = _xlsx([["n"]] + [[index] for index in range(1, 1001)])

    workbook = spreadsheets.read_workbook(raw, "xlsx", limits=spreadsheets.Limits(rows=100, cols=5))
    sheet = workbook.sheets[0]

    assert len(sheet.rows) == 100
    assert sheet.total_rows == 1001
    assert sheet.truncated_rows is True
    assert sheet.warnings == ["Only the first 100 of 1,001 rows were included."]
    assert spreadsheets.read_warnings(workbook) == ["Only the first 100 of 1,001 rows were included."]


def test_hidden_sheets_are_skipped_and_extra_sheets_are_announced() -> None:
    workbook = Workbook()
    workbook.active.title = "One"
    workbook.active.append(["a"])
    for name in ("Two", "Three"):
        workbook.create_sheet(name).append([name])
    hidden = workbook.create_sheet("Hidden")
    hidden.append(["x"])
    hidden.sheet_state = "hidden"
    raw = BytesIO()
    workbook.save(raw)

    parsed = spreadsheets.read_workbook(raw.getvalue(), "xlsx", limits=spreadsheets.Limits(sheets=2))

    assert parsed.sheet_names == ["One", "Two", "Three", "Hidden"]
    assert [sheet.name for sheet in parsed.sheets] == ["One", "Two"]
    assert parsed.warnings == ["Only the first 2 of 3 visible sheets were included."]


def test_formula_cells_without_cached_values_show_the_formula() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Revenue", 100])
    sheet.append(["Costs", 60])
    sheet["B3"] = "=B1-B2"
    raw = BytesIO()
    workbook.save(raw)

    parsed = spreadsheets.read_workbook(raw.getvalue(), "xlsx")

    assert parsed.sheets[0].rows[2] == ["", "=B1-B2"]
    assert parsed.sheets[0].formulas == [{"cell": "B3", "formula": "=B1-B2", "cached_value": ""}]


def test_a_zip_that_is_not_a_workbook_is_not_a_spreadsheet() -> None:
    output = BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("notes.txt", "hello")

    assert spreadsheets.sniff_kind(output.getvalue(), file_name="files.zip") is None
    with pytest.raises(spreadsheets.SpreadsheetError) as exc:
        spreadsheets.read_workbook(output.getvalue(), "xlsx")
    assert exc.value.code == "parse_failed"


# --- .csv ------------------------------------------------------------------------


def test_csv_delimiters_and_encodings_are_detected() -> None:
    semicolon = b"Name;Amount\nRent;4500\n"
    utf16 = "Name\tAmount\nRent\t4500\n".encode("utf-16")

    first = spreadsheets.read_workbook(semicolon, "csv", file_name="rent.csv")
    second = spreadsheets.read_workbook(utf16, "csv", file_name="rent.csv")

    assert first.delimiter == ";" and first.sheets[0].rows == [["Name", "Amount"], ["Rent", "4500"]]
    assert second.encoding == "utf-16" and second.sheets[0].rows == [["Name", "Amount"], ["Rent", "4500"]]
    assert first.sheets[0].name == "rent"


def test_a_windows_browser_calling_a_csv_an_excel_file_is_read_as_csv() -> None:
    raw = b"Date,Description,Amount\n2026-01-02,Deposit,1200.50\n"

    assert spreadsheets.sniff_kind(raw, declared_content_type="application/vnd.ms-excel", file_name="export.csv") == "csv"
    assert detect_document_content_type(raw, "application/vnd.ms-excel") == "text/csv"


def test_a_utf16_csv_export_no_longer_falls_through_to_the_declared_type() -> None:
    raw = "Date\tAmount\n2026-01-02\t1200\n".encode("utf-16")

    assert detect_document_content_type(raw, "application/vnd.ms-excel") == "text/csv"


def test_single_column_csv_is_still_a_table() -> None:
    raw = b"amount\n100\n200\n"

    workbook = spreadsheets.read_workbook(raw, "csv", file_name="amounts.csv")

    assert workbook.sheets[0].rows == [["amount"], ["100"], ["200"]]


def test_an_empty_csv_says_so() -> None:
    with pytest.raises(spreadsheets.SpreadsheetError) as exc:
        spreadsheets.read_workbook(b"\n\n", "csv")
    assert exc.value.code == "empty"


# --- legacy .xls -------------------------------------------------------------------


class _Cell:
    def __init__(self, ctype: int, value: object) -> None:
        self.ctype = ctype
        self.value = value


class _XlsSheet:
    def __init__(self, name: str, cells: list[list[_Cell]], visibility: int = 0) -> None:
        self.name = name
        self._cells = cells
        self.visibility = visibility
        self.nrows = len(cells)
        self.ncols = max((len(row) for row in cells), default=0)

    def cell(self, row: int, col: int) -> _Cell:
        line = self._cells[row]
        return line[col] if col < len(line) else _Cell(0, "")


class _XlsBook:
    datemode = 0

    def __init__(self, sheets: list[_XlsSheet]) -> None:
        self._sheets = sheets

    def sheets(self) -> list[_XlsSheet]:
        return self._sheets


def test_legacy_xls_cells_become_display_text(monkeypatch: pytest.MonkeyPatch) -> None:
    import xlrd

    book = _XlsBook(
        [
            _XlsSheet(
                "Debts",
                [
                    [_Cell(xlrd.XL_CELL_TEXT, "Lender"), _Cell(xlrd.XL_CELL_TEXT, "Balance"), _Cell(xlrd.XL_CELL_TEXT, "Due")],
                    [_Cell(xlrd.XL_CELL_TEXT, "Chase"), _Cell(xlrd.XL_CELL_NUMBER, 125000.0), _Cell(xlrd.XL_CELL_DATE, 46023.0)],
                    [_Cell(xlrd.XL_CELL_EMPTY, ""), _Cell(xlrd.XL_CELL_EMPTY, "")],
                    [_Cell(xlrd.XL_CELL_TEXT, "Paid"), _Cell(xlrd.XL_CELL_BOOLEAN, 1), _Cell(xlrd.XL_CELL_ERROR, 0)],
                ],
            ),
            _XlsSheet("Hidden", [[_Cell(xlrd.XL_CELL_TEXT, "nope")]], visibility=1),
        ]
    )
    monkeypatch.setattr(xlrd, "open_workbook", lambda **_kwargs: book)
    monkeypatch.setattr(spreadsheets, "_ole_is_encrypted_package", lambda _raw: False)

    workbook = spreadsheets.read_workbook(b"\xd0\xcf\x11\xe0 fake", "xls")

    assert workbook.kind == "xls"
    assert workbook.sheet_names == ["Debts", "Hidden"]
    assert [sheet.name for sheet in workbook.sheets] == ["Debts"]
    assert workbook.sheets[0].rows == [
        ["Lender", "Balance", "Due"],
        ["Chase", "125000", str(date(2026, 1, 1))],
        ["Paid", "True", "#ERROR"],
    ]


@pytest.mark.parametrize(
    ("message", "code"),
    [
        ("Workbook is encrypted", "locked"),
        ("Can't find workbook in OLE2 compound document", "not_a_workbook"),
        ("Expected BOF record; found b'...'", "not_a_workbook"),
        ("something else went wrong", "parse_failed"),
    ],
)
def test_legacy_xls_failures_are_named(monkeypatch: pytest.MonkeyPatch, message: str, code: str) -> None:
    import xlrd

    def _raise(**_kwargs):
        raise xlrd.XLRDError(message)

    monkeypatch.setattr(xlrd, "open_workbook", _raise)
    monkeypatch.setattr(spreadsheets, "_ole_is_encrypted_package", lambda _raw: False)

    with pytest.raises(spreadsheets.SpreadsheetError) as exc:
        spreadsheets.read_workbook(b"\xd0\xcf\x11\xe0 fake", "xls")
    assert exc.value.code == code


def test_an_encrypted_ooxml_package_is_reported_locked(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(spreadsheets, "_ole_is_encrypted_package", lambda _raw: True)

    with pytest.raises(spreadsheets.SpreadsheetError) as exc:
        spreadsheets.read_workbook(b"\xd0\xcf\x11\xe0 fake", "xls")
    assert exc.value.code == "locked"


# --- how the AI attachment uses it -------------------------------------------------


def _file(name: str):
    return SimpleNamespace(id="11111111-1111-1111-1111-111111111111", file_name=name)


def test_the_ai_reads_legacy_xls_instead_of_refusing_it(monkeypatch: pytest.MonkeyPatch) -> None:
    import xlrd

    book = _XlsBook([_XlsSheet("P&L", [[_Cell(xlrd.XL_CELL_TEXT, "Revenue"), _Cell(xlrd.XL_CELL_NUMBER, 900.0)]])])
    monkeypatch.setattr(xlrd, "open_workbook", lambda **_kwargs: book)
    monkeypatch.setattr(spreadsheets, "_ole_is_encrypted_package", lambda _raw: False)
    content: list[dict] = []
    skipped: list[dict] = []

    added, _pages, chars = _append_review_file_content(
        content=content,
        skipped=skipped,
        blocked_files=[],
        file=_file("financials.xls"),
        raw=b"\xd0\xcf\x11\xe0 fake",
        content_type="application/octet-stream",
        attached_pdf_pages=0,
        spreadsheet_text_chars=0,
    )

    assert added is True
    assert chars > 0
    assert content[0]["text"].startswith("Legacy Excel file")
    assert '"type": "xls_workbook"' in content[0]["text"]
    assert skipped == []


def test_a_locked_legacy_xls_is_a_replacement_request(monkeypatch: pytest.MonkeyPatch) -> None:
    import xlrd

    def _raise(**_kwargs):
        raise xlrd.XLRDError("Workbook is encrypted")

    monkeypatch.setattr(xlrd, "open_workbook", _raise)
    monkeypatch.setattr(spreadsheets, "_ole_is_encrypted_package", lambda _raw: False)

    extracted, skip = _extract_xls_text(_file("locked.xls"), b"\xd0\xcf\x11\xe0 fake")

    assert extracted is None
    assert skip[0] == "spreadsheet_locked"


def test_a_partially_read_workbook_is_attached_with_a_read_warning() -> None:
    raw = _xlsx([["n"]] + [[index] for index in range(1, 601)])
    content: list[dict] = []
    skipped: list[dict] = []

    added, _pages, _chars = _append_review_file_content(
        content=content,
        skipped=skipped,
        blocked_files=[],
        file=_file("big.xlsx"),
        raw=raw,
        content_type="application/octet-stream",
        attached_pdf_pages=0,
        spreadsheet_text_chars=0,
    )

    assert added is True
    assert [item["reason"] for item in skipped] == ["spreadsheet_partially_read"]
    assert "Only the first 500 of 601 rows were included." in skipped[0]["explanation"]


def test_the_empty_looking_workbook_from_production_is_now_read() -> None:
    raw = _without_dimension(_xlsx([["Profit and Loss"], ["Revenue", 48210], ["Expenses", 31000]], title="Profit and Loss"))
    content: list[dict] = []
    skipped: list[dict] = []

    added, _pages, _chars = _append_review_file_content(
        content=content,
        skipped=skipped,
        blocked_files=[],
        file=_file("EMBR_Home_QC_P_and_L_Balance_Sheet.xlsx"),
        raw=raw,
        content_type="application/octet-stream",
        attached_pdf_pages=0,
        spreadsheet_text_chars=0,
    )

    assert added is True
    assert "Sheet appears empty" not in content[0]["text"]
    assert '"Revenue", "48210"' in content[0]["text"]


def test_old_unreadable_spreadsheet_skips_are_read_again_but_other_skips_stay_cached() -> None:
    legacy = SimpleNamespace(status="skipped", skip_reason="legacy_spreadsheet_unsupported")
    locked_pdf = SimpleNamespace(status="skipped", skip_reason="password_protected")
    completed = SimpleNamespace(status="completed", skip_reason=None)

    assert _spreadsheet_skip_still_applies(legacy, content_type="application/vnd.ms-excel", file_name="old.xls") is False
    assert _spreadsheet_skip_still_applies(legacy, content_type="application/pdf", file_name="scan.pdf") is True
    assert _spreadsheet_skip_still_applies(locked_pdf, content_type="application/pdf", file_name="scan.pdf") is True
    assert _spreadsheet_skip_still_applies(completed, content_type="text/csv", file_name="a.csv") is True


# --- the preview payload ------------------------------------------------------------


def test_preview_payload_carries_sheets_rows_and_truncation() -> None:
    raw = _xlsx([["a", "b"], [1, 2]])

    payload = spreadsheets.preview_payload(spreadsheets.read_workbook(raw, "xlsx", limits=spreadsheets.PREVIEW_LIMITS))

    assert payload["kind"] == "xlsx"
    assert payload["sheets"][0]["rows"] == [["a", "b"], ["1", "2"]]
    assert payload["sheets"][0]["truncated_rows"] is False
    assert payload["sheet_names"] == ["Sheet1"]


def test_is_spreadsheet_name_is_a_cheap_gate_on_metadata() -> None:
    assert spreadsheets.is_spreadsheet_name("Debt schedule.XLSX", "application/octet-stream")
    assert spreadsheets.is_spreadsheet_name("bank.csv", None)
    assert spreadsheets.is_spreadsheet_name("export.bin", "application/vnd.ms-excel")
    assert not spreadsheets.is_spreadsheet_name("statement.pdf", "application/pdf")


# --- review follow-ups -------------------------------------------------------------


def _with_dimension(xlsx: bytes, ref: str) -> bytes:
    """A workbook whose sheet XML carries a stale <dimension> record."""

    source = zipfile.ZipFile(BytesIO(xlsx))
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data, count = re.subn(rb"<dimension ref=\"[^\"]*\"\s*/>", f'<dimension ref="{ref}"/>'.encode(), data)
                assert count == 1
            target.writestr(item, data)
    return output.getvalue()


@pytest.mark.parametrize("ref", ["A1", "A1:B2"])
def test_a_stale_dimension_record_does_not_cut_the_sheet(ref: str) -> None:
    raw = _with_dimension(_xlsx([["Account", "Amount", "Note"], ["Rent", 4500, "monthly"], ["Payroll", 12000, "biweekly"]]), ref)

    sheet = spreadsheets.read_workbook(raw, "xlsx").sheets[0]

    assert sheet.rows == [["Account", "Amount", "Note"], ["Rent", "4500", "monthly"], ["Payroll", "12000", "biweekly"]]
    assert (sheet.total_rows, sheet.total_cols) == (3, 3)
    assert sheet.warnings == []


def test_a_chart_sheet_is_listed_but_not_read_and_the_data_still_is() -> None:
    from openpyxl.chart import BarChart, Reference

    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Month", "Sales"])
    data.append(["Jan", 100])
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    chart_sheet = workbook.create_chartsheet("Chart")
    chart_sheet.add_chart(chart)
    raw = BytesIO()
    workbook.save(raw)

    parsed = spreadsheets.read_workbook(raw.getvalue(), "xlsx")

    assert parsed.sheet_names == ["Data", "Chart"]
    assert [sheet.name for sheet in parsed.sheets] == ["Data"]
    assert parsed.sheets[0].rows == [["Month", "Sales"], ["Jan", "100"]]


def test_a_count_past_the_ceiling_is_reported_as_a_floor(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(spreadsheets, "_COUNT_CEILING", 10)
    raw = _xlsx([["n"]] + [[index] for index in range(1, 30)])

    sheet = spreadsheets.read_workbook(raw, "xlsx", limits=spreadsheets.Limits(rows=5)).sheets[0]

    assert sheet.count_capped is True
    assert sheet.total_rows == 10
    assert sheet.warnings == ["Only the first 5 of more than 10 rows were included."]
    assert spreadsheets.preview_payload(spreadsheets.Workbook(kind="xlsx", sheet_names=["a"], sheets=[sheet]))["sheets"][0]["count_capped"] is True


def test_a_csv_that_starts_with_the_letters_pk_is_still_a_csv() -> None:
    raw = b"PK,Name,Amount\n1,Rent,500\n"

    assert spreadsheets.sniff_kind(raw, declared_content_type="text/csv", file_name="pk.csv") == "csv"
    assert spreadsheets.read_workbook(raw, file_name="pk.csv").sheets[0].rows[0] == ["PK", "Name", "Amount"]


def test_an_array_formula_shows_its_text_not_an_object() -> None:
    from openpyxl.worksheet.formula import ArrayFormula

    workbook = Workbook()
    sheet = workbook.active
    sheet.append([1, 2])
    sheet["C1"] = ArrayFormula("C1:C1", "=SUM(A1:B1)")
    raw = BytesIO()
    workbook.save(raw)

    parsed = spreadsheets.read_workbook(raw.getvalue(), "xlsx")

    assert parsed.sheets[0].rows[0] == ["1", "2", "=SUM(A1:B1)"]
    assert "ArrayFormula" not in json_dumps(parsed.sheets[0].rows)


def json_dumps(value: object) -> str:
    import json

    return json.dumps(value, default=str)


def test_an_html_export_declared_as_excel_is_not_a_csv() -> None:
    raw = b"<html><body><table><tr><td>Revenue</td><td>100</td></tr></table></body></html>"

    assert spreadsheets.sniff_kind(raw, declared_content_type="application/vnd.ms-excel", file_name="export.xls") is None
    with pytest.raises(spreadsheets.SpreadsheetError) as exc:
        spreadsheets.read_workbook(raw, declared_content_type="application/vnd.ms-excel", file_name="export.xls")
    assert exc.value.code == "not_a_workbook"


def test_a_utf8_sample_cut_inside_a_character_still_sniffs_as_csv() -> None:
    line = "José,Payee,1200\n".encode()
    raw = b"Name,Who,Amount\n" + line * (8192 // len(line) + 2)
    # make the 8192nd byte land inside the two-byte "é"
    cut = raw[:8192]
    while True:
        try:
            cut.decode("utf-8")
        except UnicodeDecodeError:
            break
        raw = b"x" + raw
        cut = raw[:8192]

    assert detect_document_content_type(raw, "application/vnd.ms-excel") == "text/csv"


def test_a_cp1252_csv_declared_as_excel_is_read_as_csv_not_sent_to_xlrd() -> None:
    raw = "Name,Amount\nJosé,1200\n".encode("cp1252")
    content: list[dict] = []
    skipped: list[dict] = []

    added, _pages, _chars = _append_review_file_content(
        content=content,
        skipped=skipped,
        blocked_files=[],
        file=_file("export.csv"),
        raw=raw,
        content_type="application/vnd.ms-excel",
        attached_pdf_pages=0,
        spreadsheet_text_chars=0,
    )

    assert added is True
    assert content[0]["text"].startswith("CSV file")
    assert "Jos" in content[0]["text"]
    assert skipped == []


def test_pdf_bytes_under_a_csv_name_are_attached_as_a_pdf() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    output = BytesIO()
    writer.write(output)
    content: list[dict] = []
    skipped: list[dict] = []

    added, pages, _chars = _append_review_file_content(
        content=content,
        skipped=skipped,
        blocked_files=[],
        file=_file("statement.csv"),
        raw=output.getvalue(),
        content_type="text/csv",
        attached_pdf_pages=0,
        spreadsheet_text_chars=0,
    )

    assert added is True
    assert pages == 1
    assert content[0]["text"].startswith("PDF file")
    assert not any(block.get("text", "").startswith("CSV file") for block in content)


def test_a_skip_written_by_the_current_reader_is_not_reopened() -> None:
    old = SimpleNamespace(status="skipped", skip_reason="spreadsheet_parse_failed", analysis=None)
    current = SimpleNamespace(
        status="skipped", skip_reason="spreadsheet_parse_failed", analysis={"spreadsheet_reader": spreadsheets.READER_VERSION}
    )

    assert _spreadsheet_skip_still_applies(old, content_type="application/octet-stream", file_name="broken.xlsx") is False
    assert _spreadsheet_skip_still_applies(current, content_type="application/octet-stream", file_name="broken.xlsx") is True


def test_the_preview_reads_with_the_ai_bounds() -> None:
    from app.services import bucket_ai

    assert spreadsheets.PREVIEW_LIMITS.sheets == bucket_ai._SPREADSHEET_LIMITS.sheets == bucket_ai.MAX_SPREADSHEET_SHEETS
    assert spreadsheets.PREVIEW_LIMITS.rows == bucket_ai._SPREADSHEET_LIMITS.rows == bucket_ai.MAX_SPREADSHEET_ROWS
    assert spreadsheets.PREVIEW_LIMITS.cols == bucket_ai._SPREADSHEET_LIMITS.cols == bucket_ai.MAX_SPREADSHEET_COLS
    assert spreadsheets.PREVIEW_LIMITS.cell_chars == bucket_ai.MAX_CELL_CHARS
    assert spreadsheets.PREVIEW_LIMITS.formulas == 0
    assert bucket_ai.MAX_DOCX_TEXT_CHARS == 24000
