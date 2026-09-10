"""The four downloadable templates: generated from the schemas, with every
defect of the originals fixed, and served free by the public router."""

from __future__ import annotations

import asyncio
import re
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.services import business_statement_schema as bss
from app.services import financial_templates_xlsx as templates
from app.services.bucket_ai import MAX_SPREADSHEET_ROWS, MAX_SPREADSHEET_SHEETS
from app.services.bucket_evidence import filename_evidence_classification
from app.services.financial_statements import DEBT_COLUMN_LABELS, DEBT_COLUMNS

KINDS = tuple(templates.SLUGS.values())


def _open(kind):
    return load_workbook(BytesIO(templates.build_workbook(kind)))


def _named_cell(wb, name):
    """The cell a defined name points at."""
    target = wb.defined_names[name].attr_text
    sheet, ref = target.rsplit("!", 1)
    return wb[sheet.strip("'")][ref.replace("$", "")]


def _cells(ws):
    return [cell for row in ws.iter_rows() for cell in row]


@pytest.mark.parametrize("kind", KINDS)
def test_every_workbook_opens_and_stays_under_the_analyzers_budget(kind):
    wb = _open(kind)
    assert len(wb.sheetnames) == 1 <= MAX_SPREADSHEET_SHEETS
    assert wb.active.max_row <= MAX_SPREADSHEET_ROWS
    assert wb.active.print_area


@pytest.mark.parametrize("kind", KINDS)
def test_every_schema_key_has_a_defined_name_and_a_hidden_key_beside_it(kind):
    wb = _open(kind)
    prefix = templates.NAME_PREFIX[kind]
    for key in templates.schema_keys(kind):
        assert f"{prefix}.{key}" in wb.defined_names, key
    ws = wb.active
    hidden = [letter for letter, dim in ws.column_dimensions.items() if dim.hidden]
    assert hidden, "no hidden key column"
    if kind != "debt_schedule":
        keys = {cell.value for cell in ws["C"] if cell.value}
        assert set(templates.schema_keys(kind)) <= keys


@pytest.mark.parametrize("kind", KINDS)
def test_no_year_is_written_into_any_cell(kind):
    for cell in _cells(_open(kind).active):
        assert not (isinstance(cell.value, str) and re.search(r"20\d{2}", cell.value)), cell.coordinate


@pytest.mark.parametrize(
    ("kind", "totals"),
    [
        ("p_and_l", ["gross_profit", "total_operating_expenses", "operating_income", "net_income", "ebitda", "owner_compensation"]),
        ("balance_sheet", ["total_current_assets", "total_fixed_assets", "total_other_assets", "total_assets", "total_current_liabilities", "total_long_term_liabilities", "total_liabilities", "total_equity", "total_liabilities_and_equity", "imbalance"]),
        ("pfs", ["total_assets", "total_liabilities", "net_worth", "total_income", "total_contingent"]),
        ("debt_schedule", ["total_balance", "total_monthly_payment"]),
    ],
)
def test_every_total_is_a_formula_on_a_locked_cell(kind, totals):
    wb = _open(kind)
    assert wb.active.protection.sheet is True
    for key in totals:
        cell = _named_cell(wb, f"{templates.NAME_PREFIX[kind]}.{key}")
        assert isinstance(cell.value, str) and cell.value.startswith("="), key
        assert cell.protection.locked is True, key
        assert cell.number_format == templates.MONEY_FORMAT, key


@pytest.mark.parametrize("kind", KINDS)
def test_inputs_are_unlocked_and_money_inputs_share_one_format(kind):
    wb = _open(kind)
    prefix = templates.NAME_PREFIX[kind]
    for key in templates.schema_keys(kind):
        cell = _named_cell(wb, f"{prefix}.{key}")
        if isinstance(cell.value, str) and cell.value.startswith("="):
            continue   # a subtotal
        assert cell.protection.locked is False, key
        assert cell.value is None, key
    if kind in bss.KINDS:
        for section in bss.SCHEMA_FOR[kind].sections:
            for row in section.rows:
                assert _named_cell(wb, f"{prefix}.{row.key}").number_format == templates.MONEY_FORMAT


def test_the_p_and_l_ebitda_memo_is_built_from_the_addback_rows():
    wb = _open("p_and_l")
    ebitda = _named_cell(wb, "pl.ebitda")
    net_income = _named_cell(wb, "pl.net_income")
    assert net_income.coordinate in ebitda.value
    for key in ("interest", "income_taxes", "depreciation_and_amortization"):
        memo = _named_cell(wb, f"pl.memo_{key}")
        assert memo.value == f"={_named_cell(wb, f'pl.{key}').coordinate}"
        assert memo.coordinate in ebitda.value
    assert "pl.memo_taxes_and_licenses" not in wb.defined_names
    # Dates are blank cells with a comment, in a date format, never a title.
    start = _named_cell(wb, "pl.period_start")
    assert start.value is None and start.comment is not None and start.number_format == templates.DATE_FORMAT


def test_the_balance_sheet_implies_equity_when_blank_and_prints_the_difference():
    wb = _open("balance_sheet")
    equity = _named_cell(wb, "bs.total_equity").value
    assert equity.startswith("=IF(COUNT(")
    assert _named_cell(wb, "bs.total_assets").coordinate in equity
    imbalance = _named_cell(wb, "bs.imbalance").value
    assert _named_cell(wb, "bs.total_liabilities_and_equity").coordinate in imbalance


def test_the_debt_schedule_header_is_the_column_labels_with_keys_beneath():
    wb = _open("debt_schedule")
    ws = wb.active
    header_row = next(row for row in range(1, 10) if ws.cell(row=row, column=1).value == "Lender")
    assert tuple(ws.cell(row=header_row, column=i).value for i in range(1, len(DEBT_COLUMNS) + 1)) == DEBT_COLUMN_LABELS
    assert tuple(ws.cell(row=header_row + 1, column=i).value for i in range(1, len(DEBT_COLUMNS) + 1)) == DEBT_COLUMNS
    assert ws.row_dimensions[header_row + 1].hidden
    assert _named_cell(wb, "ds.r7.lender").row == header_row + 1 + 7
    assert _named_cell(wb, "ds.r15.notes").protection.locked is False


def test_the_pfs_sheet_says_the_schedules_are_on_screen():
    texts = [cell.value for cell in _cells(_open("pfs").active) if isinstance(cell.value, str)]
    assert any("completed on screen" in text for text in texts)
    assert any("No Social Security Number" in text for text in texts)


def test_the_bytes_are_deterministic():
    first = templates.build_workbook.__wrapped__("p_and_l")
    second = templates.build_workbook.__wrapped__("p_and_l")
    assert first == second
    assert templates.build_workbook("p_and_l") is templates.build_workbook("p_and_l")   # cached


def test_the_attachment_filenames_route_a_filled_copy_to_its_checklist_row():
    names = templates.ATTACHMENT_FILENAMES
    assert filename_evidence_classification(names["p_and_l"]) == "current_p_and_l"
    # "balance_sheet" once the evidence vocabulary learns the word; a P&L until then.
    assert filename_evidence_classification(names["balance_sheet"]) in {"balance_sheet", "current_p_and_l"}
    assert filename_evidence_classification(names["debt_schedule"]) == "debt_schedule"
    assert filename_evidence_classification(names["pfs"]) == "personal_financial_statement"


def test_the_blank_p_and_l_reaches_the_model_with_every_row_label():
    from app.services.bucket_ai import _extract_xlsx_text

    text, skip = _extract_xlsx_text(SimpleNamespace(id="f", file_name="x.xlsx"), templates.build_workbook("p_and_l"))
    assert skip is None
    for section in bss.PL_SECTIONS:
        assert section.label in text
        for row in section.rows:
            assert row.label in text, row.label
    assert "EBITDA (memo)" in text


def test_the_public_route_serves_an_attachment_and_404s_an_unknown_slug():
    from app.routers.public import financial_template_download

    response = asyncio.run(financial_template_download("profit-and-loss"))
    assert response.media_type == templates.MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="Qualified Commercial - Profit and Loss Statement.xlsx"'
    assert response.headers["cache-control"] == "public, max-age=86400"
    assert response.body == templates.build_workbook("p_and_l")
    with pytest.raises(HTTPException) as caught:
        asyncio.run(financial_template_download("tax-return"))
    assert caught.value.status_code == 404
    assert templates.workbook_for_slug("nope") is None


# --- The fifth download: the four forms as four tabs of one workbook. --------
#
# The owner asked for "a google sheet type of interface with all these already
# in one single sheet that we can forward the client or their accountant". The
# risk this section guards is that _Form is shared with the four singles: it
# now picks wb.active for the first form and a new sheet for each one after,
# so a bug there would corrupt all five downloads at once.


def _packet():
    return load_workbook(BytesIO(templates.build_packet_workbook()))


@pytest.mark.parametrize("kind", KINDS)
def test_each_single_workbook_is_still_exactly_one_sheet(kind):
    """The packet must not leak extra tabs into the four singles."""
    assert _open(kind).sheetnames == [templates.SHEET_TITLES[kind]]


def test_the_packet_is_the_four_forms_as_tabs_in_the_order_a_reader_expects():
    wb = _packet()
    assert wb.sheetnames == [
        "Profit and Loss",
        "Balance Sheet",
        "Business Debt Schedule",
        "Personal Financial Statement",
    ]
    assert [templates.SHEET_TITLES[kind] for kind in templates.PACKET_KINDS] == wb.sheetnames
    assert wb.properties.title == "Financial Package"
    for ws in wb.worksheets:
        assert ws.print_area, ws.title
        assert ws.protection.sheet is True, ws.title


def test_the_packet_stays_inside_the_analyzers_sheet_and_row_budget():
    """A filled copy uploaded back to a room must be read whole, not truncated."""
    wb = _packet()
    assert len(wb.sheetnames) <= MAX_SPREADSHEET_SHEETS
    for ws in wb.worksheets:
        assert ws.max_row <= MAX_SPREADSHEET_ROWS, (ws.title, ws.max_row)


@pytest.mark.parametrize("kind", KINDS)
def test_every_defined_name_of_every_single_workbook_is_in_the_packet_on_its_own_tab(kind):
    """The pl. / bs. / ds. / pfs. prefixes keep the four sets apart in one
    workbook, and each name still carries its own sheet title, so a reader
    written against a single template reads the packet unchanged."""
    single = _open(kind)
    packet = _packet()
    assert single.defined_names, kind
    for name, defined in single.defined_names.items():
        assert name in packet.defined_names, name
        assert packet.defined_names[name].attr_text == defined.attr_text, name
        assert templates.SHEET_TITLES[kind] in defined.attr_text, name
        assert _named_cell(packet, name).parent.title == templates.SHEET_TITLES[kind], name


@pytest.mark.parametrize("kind", KINDS)
def test_each_packet_tab_is_cell_for_cell_the_single_workbook(kind):
    single = _open(kind).active
    tab = _packet()[templates.SHEET_TITLES[kind]]
    assert [cell.value for cell in _cells(tab)] == [cell.value for cell in _cells(single)]
    assert tab.max_row == single.max_row and tab.max_column == single.max_column


def test_the_packet_filename_is_not_read_as_a_bank_statement():
    """`filename_evidence_classification` classifies anything with "statement"
    in the name it cannot place more precisely as a bank_statement. The packet
    is four documents at once and belongs to no single checklist row, so its
    name deliberately avoids that word and classifies as None — otherwise a
    filled copy uploaded back would be filed against bank statements."""
    name = templates.PACKET_ATTACHMENT_FILENAME
    assert name == "Qualified Commercial - Financial Package.xlsx"
    assert "statement" not in name.casefold()
    assert filename_evidence_classification(name) is None


def test_the_packet_bytes_are_deterministic_and_built_once():
    assert templates.build_packet_workbook.__wrapped__() == templates.build_packet_workbook.__wrapped__()
    assert templates.build_packet_workbook() is templates.build_packet_workbook()


def test_the_public_route_serves_the_packet_like_the_other_four():
    from app.routers.public import financial_template_download

    assert templates.PACKET_SLUG == "financial-package"
    assert templates.PACKET_SLUG not in templates.SLUGS   # not a schema kind
    response = asyncio.run(financial_template_download(templates.PACKET_SLUG))
    assert response.media_type == templates.MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="Qualified Commercial - Financial Package.xlsx"'
    )
    assert response.headers["cache-control"] == "public, max-age=86400"
    assert response.body == templates.build_packet_workbook()
