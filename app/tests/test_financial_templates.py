from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.application_programs import (
    _coverage_for_files,
    _filename_suggestions,
    _matching_evidence,
    _matching_evidence_files,
)
from app.services.financial_templates import TEMPLATE_VERSION, template_for_requirement


def _requirement(key: str, label: str):
    return SimpleNamespace(requirement_key=key, label=label, category="financials")


def _file(name: str, *, requested_document_id=None):
    return SimpleNamespace(
        id=uuid4(),
        file_name=name,
        requested_document_id=requested_document_id,
        created_at=datetime.now(UTC),
        statement_period=None,
    )


def _analysis(classification: str):
    return SimpleNamespace(
        id=uuid4(),
        status="completed",
        classification=classification,
        analysis={},
    )


def test_financial_templates_are_versioned_and_readable() -> None:
    pfs = template_for_requirement("owner_personal_financial_statement")
    debt = template_for_requirement("business_debt_schedule")
    statements = template_for_requirement("ytd_p_and_l_balance_sheet")

    assert pfs and pfs.content.startswith(b"%PDF") and TEMPLATE_VERSION in pfs.filename
    assert debt and debt.content.startswith(b"PK") and TEMPLATE_VERSION in debt.filename
    assert statements and statements.content.startswith(b"PK") and TEMPLATE_VERSION in statements.filename
    assert len(PdfReader(BytesIO(pfs.content)).pages) == 1

    debt_book = load_workbook(BytesIO(debt.content), data_only=False)
    assert debt_book.sheetnames == ["Debt Schedule"]
    assert debt_book["Debt Schedule"]["D29"].value == "=SUM(D4:D28)"
    assert len(debt_book["Debt Schedule"].data_validations.dataValidation) == 1

    statement_book = load_workbook(BytesIO(statements.content), data_only=False)
    assert statement_book.sheetnames == ["Profit and Loss", "Balance Sheet"]
    assert statement_book["Profit and Loss"]["B17"].value == "=B14-B15-B16"
    assert statement_book["Balance Sheet"]["B10"].value == "=SUM(B4:B9)"


def test_combined_financial_requirement_needs_both_statements() -> None:
    requirement = _requirement("ytd_p_and_l_balance_sheet", "YTD P&L and balance sheet")
    pnl = _file("profit-and-loss.pdf")
    balance = _file("balance-sheet.pdf")

    selected, complete, coverage = _matching_evidence(
        requirement,
        None,
        [pnl],
        {pnl.id: _analysis("profit_and_loss")},
    )
    assert selected is pnl
    assert complete is False
    assert coverage["profit_and_loss"] is True
    assert coverage["balance_sheet"] is False

    selected, complete, coverage = _matching_evidence(
        requirement,
        None,
        [pnl, balance],
        {
            pnl.id: _analysis("profit_and_loss"),
            balance.id: _analysis("balance_sheet"),
        },
    )
    assert selected is pnl or selected is balance
    assert complete is True
    assert coverage["balance_sheet"] is True

    matches, complete, coverage = _matching_evidence_files(
        requirement,
        None,
        [pnl, balance],
        {
            pnl.id: _analysis("profit_and_loss"),
            balance.id: _analysis("balance_sheet"),
        },
    )
    assert {file.id for file in matches} == {pnl.id, balance.id}
    assert complete is True
    assert coverage["current"] == coverage["required"] == 2


def test_bank_and_tax_coverage_count_distinct_periods_across_documents() -> None:
    bank_requirement = _requirement(
        "business_bank_statements_6_months", "Last 6 months business bank statements"
    )
    bank_files = [_file(f"statement-2026-{month:02d}.pdf") for month in range(1, 7)]
    complete, coverage = _coverage_for_files(bank_requirement, bank_files, {})

    assert complete is True
    assert coverage["months"] == [f"2026-{month:02d}" for month in range(1, 7)]
    assert coverage["current"] == coverage["required"] == 6

    tax_requirement = _requirement(
        "business_tax_returns_2_years", "Last 2 years business tax returns"
    )
    tax_files = [_file("business-tax-return-2024.pdf"), _file("business-tax-return-2025.pdf")]
    complete, coverage = _coverage_for_files(tax_requirement, tax_files, {})

    assert complete is True
    assert coverage["years"] == ["2024", "2025"]
    assert coverage["current"] == coverage["required"] == 2


def test_high_signal_legacy_upload_is_suggested_but_not_content_verified() -> None:
    requirement = _requirement("business_debt_schedule", "Business debt schedule")
    legacy = _file("Business Debt Schedule.xlsx")

    selected, complete, _coverage = _matching_evidence(requirement, None, [legacy], {})
    suggestions = _filename_suggestions(requirement, [legacy], {})

    assert selected is None
    assert complete is False
    assert suggestions == [legacy]


def test_staff_linked_evidence_remains_the_selected_file() -> None:
    requirement = _requirement("business_debt_schedule", "Business debt schedule")
    staff_choice = _file("supporting-data.pdf")
    newer_classified = _file("debt-schedule.pdf")
    newer_classified.created_at = datetime.now(UTC).replace(microsecond=999999)

    selected, _, _ = _matching_evidence(
        requirement,
        None,
        [staff_choice, newer_classified],
        {newer_classified.id: _analysis("debt_schedule")},
        preferred_file_id=staff_choice.id,
        trust_preferred=True,
    )
    assert selected is staff_choice


def test_unlinked_filename_alone_does_not_satisfy_a_requirement() -> None:
    requirement = _requirement("business_debt_schedule", "Business debt schedule")
    misleading_name = _file("business-debt-schedule.xlsx")

    selected, complete, coverage = _matching_evidence(
        requirement,
        None,
        [misleading_name],
        {},
    )

    assert selected is None
    assert complete is False
    assert coverage["matched_files"] == 0
