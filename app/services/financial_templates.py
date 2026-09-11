"""Versioned blank financial templates used by guarded intake chat actions."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

TEMPLATE_VERSION = "2026.1"


@dataclass(frozen=True)
class FinancialTemplate:
    kind: str
    filename: str
    content_type: str
    content: bytes


def template_for_requirement(requirement_key: str) -> FinancialTemplate | None:
    if requirement_key == "owner_personal_financial_statement":
        return _personal_financial_statement()
    if requirement_key == "business_debt_schedule":
        return _debt_schedule()
    if requirement_key == "ytd_p_and_l_balance_sheet":
        return _pnl_balance_sheet()
    return None


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _minimal_pdf(lines: list[str]) -> bytes:
    """Render a portable text PDF without requiring system Cairo/Pango libraries."""
    chunks = [lines[index : index + 46] for index in range(0, len(lines), 46)] or [[""]]
    objects: list[bytes] = []

    def add_object(body: str) -> int:
        objects.append(body.encode("latin-1", "replace"))
        return len(objects)

    catalog_id = add_object("<< /Type /Catalog /Pages 2 0 R >>")
    pages_id = add_object("<< /Type /Pages /Kids [] /Count 0 >>")
    font_id = add_object("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    bold_font_id = add_object("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")
    page_ids: list[int] = []
    for chunk in chunks:
        text_ops = ["BT /F1 9 Tf 42 756 Td 14 TL"]
        for index, line in enumerate(chunk):
            font = "/F2 14 Tf" if index == 0 else "/F1 9 Tf"
            text_ops.append(f"{font} ({_pdf_escape(line)}) Tj T*")
        text_ops.append("ET")
        stream = "\n".join(text_ops)
        content_id = add_object(
            f"<< /Length {len(stream.encode('latin-1', 'replace'))} >>\n"
            f"stream\n{stream}\nendstream"
        )
        page_id = add_object(
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R /F2 {bold_font_id} 0 R >> >> "
            f"/Contents {content_id} 0 R >>"
        )
        page_ids.append(page_id)
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[pages_id - 1] = (
        f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("latin-1")
    )
    objects[catalog_id - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode("latin-1")

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("latin-1"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_at = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("latin-1"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n".encode("latin-1")
    )
    return bytes(output)


def _personal_financial_statement() -> FinancialTemplate:
    blank = "_" * 24
    lines = [
        "Personal Financial Statement",
        f"Qualified Commercial blank template - Version {TEMPLATE_VERSION}",
        "",
        f"Name: {blank}    Statement date: {blank}",
        f"Address: {blank}    Business: {blank}",
        "",
        "ASSETS                                      AMOUNT",
        f"Cash and checking accounts                 {blank}",
        f"Savings and marketable securities          {blank}",
        f"Retirement accounts                        {blank}",
        f"Real estate                                {blank}",
        f"Business interests                         {blank}",
        f"Vehicles and other assets                  {blank}",
        f"Other assets                               {blank}",
        f"Total assets                               {blank}",
        "",
        "LIABILITIES                                 AMOUNT",
        f"Credit cards                               {blank}",
        f"Installment loans                          {blank}",
        f"Mortgages                                  {blank}",
        f"Taxes payable                              {blank}",
        f"Business obligations                      {blank}",
        f"Other liabilities                         {blank}",
        f"Contingent liabilities                    {blank}",
        f"Total liabilities                         {blank}",
        f"Net worth                                  {blank}",
        "",
        "I certify that this information is complete and accurate to the best of my knowledge.",
        "",
        f"Signature: {blank}    Date: {blank}",
    ]
    return FinancialTemplate(
        kind="pfs_pdf",
        filename=f"QC-Personal-Financial-Statement-v{TEMPLATE_VERSION}.pdf",
        content_type="application/pdf",
        content=_minimal_pdf(lines),
    )


def _workbook_bytes(workbook) -> bytes:
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def _style_sheet(sheet, widths: dict[str, float]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet.freeze_panes = "A4"
    for cell in sheet[1]:
        cell.font = Font(size=15, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="214A9A")
    for cell in sheet[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF9")
        cell.alignment = Alignment(wrap_text=True)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = sheet.dimensions


def _debt_schedule() -> FinancialTemplate:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Debt Schedule"
    sheet.append([f"Business Debt Schedule · QC template v{TEMPLATE_VERSION}"])
    sheet.append(["Enter each current obligation. Do not include personal debt unless it guarantees the business obligation."])
    sheet.append(["Creditor", "Debt type", "Original amount", "Current balance", "Rate %", "Monthly payment", "Maturity", "Secured by", "Current?"])
    for row in range(4, 29):
        sheet.append([None] * 9)
        sheet[f"C{row}"].number_format = '$#,##0.00'
        sheet[f"D{row}"].number_format = '$#,##0.00'
        sheet[f"E{row}"].number_format = '0.00%'
        sheet[f"F{row}"].number_format = '$#,##0.00'
        sheet[f"G{row}"].number_format = 'mm/dd/yyyy'
    total_row = 29
    sheet[f"A{total_row}"] = "Totals"
    sheet[f"D{total_row}"] = f"=SUM(D4:D{total_row - 1})"
    sheet[f"F{total_row}"] = f"=SUM(F4:F{total_row - 1})"
    sheet[f"D{total_row}"].number_format = sheet[f"F{total_row}"].number_format = '$#,##0.00'
    validation = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"I4:I{total_row - 1}")
    _style_sheet(sheet, {"A": 24, "B": 20, "C": 16, "D": 16, "E": 11, "F": 17, "G": 14, "H": 24, "I": 12})
    return FinancialTemplate(
        kind="debt_schedule_xlsx",
        filename=f"QC-Business-Debt-Schedule-v{TEMPLATE_VERSION}.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=_workbook_bytes(workbook),
    )


def _pnl_balance_sheet() -> FinancialTemplate:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    pnl = workbook.active
    pnl.title = "Profit and Loss"
    pnl.append([f"Profit and Loss · QC template v{TEMPLATE_VERSION}"])
    pnl.append(["Business name", None, "Period start", None, "Period end", None])
    pnl.append(["Category", "Current period", "Prior comparable period", "Notes"])
    labels = [
        "Gross revenue",
        "Cost of goods sold",
        "Gross profit",
        "Payroll",
        "Rent / occupancy",
        "Utilities",
        "Insurance",
        "Marketing",
        "Professional fees",
        "Other operating expenses",
        "Operating income",
        "Interest expense",
        "Taxes",
        "Net income",
    ]
    for index, label in enumerate(labels, start=4):
        pnl.append([label, None, None, None])
        pnl[f"B{index}"].number_format = pnl[f"C{index}"].number_format = '$#,##0.00'
    pnl["B6"] = "=B4-B5"
    pnl["C6"] = "=C4-C5"
    pnl["B14"] = "=B6-SUM(B7:B13)"
    pnl["C14"] = "=C6-SUM(C7:C13)"
    pnl["B17"] = "=B14-B15-B16"
    pnl["C17"] = "=C14-C15-C16"
    _style_sheet(pnl, {"A": 30, "B": 19, "C": 22, "D": 36})

    balance = workbook.create_sheet("Balance Sheet")
    balance.append([f"Balance Sheet · QC template v{TEMPLATE_VERSION}"])
    balance.append(["Business name", None, "As of", None])
    balance.append(["Assets", "Amount", "Liabilities and equity", "Amount"])
    assets = ["Cash", "Accounts receivable", "Inventory", "Prepaid assets", "Fixed assets", "Other assets", "Total assets"]
    liabilities = ["Accounts payable", "Short-term debt", "Long-term debt", "Other liabilities", "Owner equity", "Retained earnings", "Total liabilities and equity"]
    for index, (asset, liability) in enumerate(zip(assets, liabilities, strict=True), start=4):
        balance.append([asset, None, liability, None])
        balance[f"B{index}"].number_format = balance[f"D{index}"].number_format = '$#,##0.00'
    balance["B10"] = "=SUM(B4:B9)"
    balance["D10"] = "=SUM(D4:D9)"
    _style_sheet(balance, {"A": 30, "B": 18, "C": 32, "D": 18})
    for sheet in (pnl, balance):
        sheet["A1"].font = Font(size=15, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="214A9A")
    return FinancialTemplate(
        kind="pnl_balance_sheet_xlsx",
        filename=f"QC-P-and-L-Balance-Sheet-v{TEMPLATE_VERSION}.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=_workbook_bytes(workbook),
    )
