"""The four financial templates the desk hands out, generated from the schemas.

The owner shared four spreadsheets — a profit and loss, a balance sheet, a
personal financial statement and a business debt schedule — and asked for them
hosted on the website. They are not committed as binaries: each is written
here, by openpyxl, from the same row definitions the on-screen forms render
(`business_statement_schema`, `pfs_schema`, `DEBT_COLUMNS`), so the spreadsheet
a borrower downloads can never drift from the form their advisor sends.

The rows themselves live in `sheet_layout`: one `Sheet` per kind, shared with
the on-screen worksheet grid, whose row numbers are the numbers this module
writes and whose `xlsx_name`s are the defined names it declares. This module
is only the renderer — styles, protection, formulas resolved from keys to
cell addresses, the hidden key column — and `test_sheet_layout.py` pins the
two together row for row.

What the generator fixes in the originals, deliberately:

- dates are blank date cells with a comment, never a title string like
  "January 1 to December 31" of some year — so a template never says the
  wrong year and the analyzer never reads a year off a heading;
- one 2-decimal accounting format everywhere, totals included;
- subtotals and totals are formulas on locked cells; the sheet is protected
  and only the input cells are unlocked, so a total cannot be typed over;
- an EBITDA memo block on the P&L from the `addback` rows, and "Total
  liabilities and equity" plus an "Unreconciled difference" line on the
  balance sheet;
- a print area, no styled empty range, every sheet well under the analyzer's
  budget (`MAX_SPREADSHEET_ROWS`, `MAX_SPREADSHEET_SHEETS`), so a filled copy
  uploaded back is read whole;
- a hidden column carrying the schema key beside every input, and a workbook
  defined name equal to it (`pl.gross_revenue`, `bs.cash_in_bank`,
  `ds.r7.lender`, `pfs.cash_on_hand`), so a deterministic reader is a
  follow-on rather than a rewrite;
- pinned document properties and zip timestamps, so the bytes are the same on
  every build and the edge can cache them by content.

There is a fifth download: the same four forms as four tabs of one workbook
(`build_packet_workbook`, slug "financial-package"), the "one sheet we can
forward the client or their accountant" the owner asked for. It is built from
the very same renderer — nothing about a form changes because it is a tab
rather than a file — so the packet cannot drift from the four singles either.

A workbook written by openpyxl carries no cached formula values. The analyzer
shows both the value view and the formula view, and its typed prompt block
tells the model to copy the lines and leave a total null when the document
does not print one; the extractor then computes totals from the lines.

Labels are English. The resource page says so on its Spanish card.
"""

from __future__ import annotations

import functools
import re
import zipfile
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.services import business_statement_schema as bss
from app.services import pfs_schema, sheet_layout
from app.services.financial_statements import DEBT_COLUMNS
from app.services.sheet_layout import (  # noqa: F401 — re-exported; callers read them here
    DEBT_ROWS,
    NAME_PREFIX,
    SHEET_TITLES,
    Cell,
    Row,
    Sheet,
)

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: Public URL slug → workbook kind. One kind per slug; the combined packet is
#: not a kind (it has no schema of its own) and lives at PACKET_SLUG instead.
SLUGS: dict[str, str] = {
    "profit-and-loss": "p_and_l",
    "balance-sheet": "balance_sheet",
    "business-debt-schedule": "debt_schedule",
    "personal-financial-statement": "pfs",
}

#: The fifth download: all four forms as four tabs of one workbook, in this
#: order. Deliberately outside SLUGS — SLUGS maps to a `business_statement_schema`
#: kind and every caller reads it that way.
PACKET_SLUG = "financial-package"
PACKET_KINDS: tuple[str, ...] = ("p_and_l", "balance_sheet", "debt_schedule", "pfs")
PACKET_TITLE = "Financial Package"

#: The packet's attachment filename. It must NOT contain the word "statement":
#: `bucket_evidence.filename_evidence_classification` falls through to a
#: `"statement" in value` branch that returns "bank_statement" for anything it
#: cannot place more precisely, so a filled copy of the combined workbook
#: uploaded back to a room would be routed to the bank-statement checklist row.
#: A packet is four documents at once and belongs to no single row, so the name
#: is chosen to classify as None and be filed by hand. Pinned by a test.
PACKET_ATTACHMENT_FILENAME = "Qualified Commercial - Financial Package.xlsx"

#: The attachment filename per kind. Each passes
#: `bucket_evidence.filename_evidence_classification`, so a filled copy
#: uploaded back to a room routes to its checklist row before analysis.
ATTACHMENT_FILENAMES: dict[str, str] = {
    "p_and_l": "Qualified Commercial - Profit and Loss Statement.xlsx",
    "balance_sheet": "Qualified Commercial - Balance Sheet.xlsx",
    "debt_schedule": "Qualified Commercial - Business Debt Schedule.xlsx",
    "pfs": "Qualified Commercial - Personal Financial Statement.xlsx",
}

MONEY_FORMAT = '#,##0.00;(#,##0.00);"-"'
DATE_FORMAT = "yyyy-mm-dd"
RATE_FORMAT = "0.00"

_AUTHOR = "Qualified Commercial"
#: Pinned so two builds are byte-identical. Not a year that appears in a cell.
_PINNED_STAMP = datetime(1980, 1, 1)

_TITLE_FONT = Font(bold=True, size=14)
_HEADING_FONT = Font(bold=True, size=11)
_BOLD = Font(bold=True)
_MUTED = Font(italic=True, color="6B7280")
_HEADING_FILL = PatternFill("solid", fgColor="F3F4F6")
_INPUT_FILL = PatternFill("solid", fgColor="FFFBEB")
_UNLOCKED = Protection(locked=False)
_RIGHT = Alignment(horizontal="right")
_WRAP = Alignment(wrap_text=True, vertical="top")

#: A formula in the layout is written against keys in braces; the renderer
#: turns each into the cell address the layout's (c, r) says.
_FORMULA_REF = re.compile(r"\{([^{}]+)\}")


def _sheet_for(wb: Workbook, kind: str, *, first: bool) -> Worksheet:
    """The sheet one form writes into.

    A workbook openpyxl creates already has one sheet, so the FIRST form built
    into a workbook takes it and every form after it adds one — which is what
    turns the four single-tab templates into the four tabs of the packet, in
    build order. Told explicitly rather than inferred from whether the active
    sheet looks empty: `wb.active` is the first sheet forever, so a builder
    that guessed would silently overwrite tab one.
    """
    ws: Worksheet = wb.active if first else wb.create_sheet()
    ws.title = SHEET_TITLES[kind]
    return ws


class _Form:
    """The renderer for one `Sheet`: a label / input / hidden-key layout on
    the form sheets, and a wide grid for the debt schedule and the PFS
    schedule blocks.

    Two key schemes, both a workbook detail the layout never sees. The form
    sheets carry the schema key in hidden column C beside every input. The
    debt schedule carries the column keys on a hidden row under its column
    head and the line's ordinal in a hidden column past the last one.
    """

    LABEL, INPUT, KEY = "A", "B", "C"

    def __init__(self, wb: Workbook, sheet: Sheet, *, first: bool = True) -> None:
        self.ws: Worksheet = _sheet_for(wb, sheet.kind, first=first)
        self.wb = wb
        self.sheet = sheet
        self.prefix = sheet.name_prefix
        self.row = 0
        self.cells: dict[str, int] = {}
        self.keys_beside_inputs = sheet.kind != "debt_schedule"
        self.key_column = 3 if self.keys_beside_inputs else len(sheet.columns) + 1
        self.last_column = max(column.c for column in sheet.columns)
        self._validations: dict[tuple[str, ...], DataValidation] = {}
        # Every keyed or named cell's address, from the layout's (c, r), so a
        # formula can be resolved before the cell it points at is written.
        self.refs: dict[str, str] = {}
        for row in sheet.rows:
            for cell in row.cells:
                address = f"{get_column_letter(cell.c)}{row.r}"
                if cell.key is not None:
                    self.refs[cell.key] = address
                if cell.xlsx_name is not None:
                    self.refs.setdefault(cell.xlsx_name, address)
        for column in sheet.columns:
            self.ws.column_dimensions[get_column_letter(column.c)].width = column.width
        self.ws.column_dimensions[get_column_letter(self.key_column)].hidden = True

    def resolve(self, formula: str) -> str:
        return _FORMULA_REF.sub(lambda match: self.refs[match.group(1)], formula)

    def band(self, colspan: int) -> list[int]:
        """The first `colspan` sheet columns — what a heading fills and a
        note merges across."""
        return [column.c for column in self.sheet.columns[:colspan]]

    def _name(self, key: str, row: int, column: int = 2) -> None:
        name = f"{self.prefix}.{key}"
        self.wb.defined_names[name] = DefinedName(
            name=name, attr_text=f"'{self.ws.title}'!${get_column_letter(column)}${row}"
        )
        self.cells[key] = row

    def _validation(self, options: tuple[str, ...]) -> DataValidation:
        """One list validation per option set per sheet."""
        found = self._validations.get(options)
        if found is None:
            found = DataValidation(
                type="list", formula1='"' + ",".join(options) + '"', allow_blank=True
            )
            self.ws.add_data_validation(found)
            self._validations[options] = found
        return found

    def line(self, text: str, font: Font) -> None:
        """A title or subtitle: one styled cell in column A."""
        self.row += 1
        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.font = font

    def heading(self, text: str, columns: list[int]) -> None:
        self.row += 1
        for column in columns:
            cell = self.ws.cell(row=self.row, column=column)
            cell.fill = _HEADING_FILL
            cell.font = _HEADING_FONT
        self.ws.cell(row=self.row, column=columns[0]).value = text

    def text_field(self, key: str, label: str, *, input: str = "text", options=None, hint: str | None = None) -> None:
        self.row += 1
        self.ws.cell(row=self.row, column=1, value=label)
        cell = self.ws.cell(row=self.row, column=2)
        cell.protection = _UNLOCKED
        cell.fill = _INPUT_FILL
        if input == "date":
            cell.number_format = DATE_FORMAT
            cell.comment = Comment(hint or "Enter as a date (year-month-day).", _AUTHOR)
        elif input == "select" and options:
            self._validation(tuple(options)).add(cell)
            cell.comment = Comment(hint or "Choose: " + " or ".join(options) + ".", _AUTHOR)
        elif hint:
            cell.comment = Comment(hint, _AUTHOR)
        if self.keys_beside_inputs:
            self.ws.cell(row=self.row, column=3, value=key)
        self._name(key, self.row)

    def money_row(self, key: str, label: str, *, hint: str | None = None) -> int:
        self.row += 1
        self.ws.cell(row=self.row, column=1, value=label)
        cell = self.ws.cell(row=self.row, column=2)
        cell.number_format = MONEY_FORMAT
        cell.protection = _UNLOCKED
        cell.fill = _INPUT_FILL
        cell.alignment = _RIGHT
        if hint:
            cell.comment = Comment(hint, _AUTHOR)
        self.ws.cell(row=self.row, column=3, value=key)
        self._name(key, self.row)
        return self.row

    def formula_row(self, key: str, label: str, formula: str, *, emphasis: bool = True) -> int:
        self.row += 1
        label_cell = self.ws.cell(row=self.row, column=1, value=label)
        cell = self.ws.cell(row=self.row, column=2, value=formula)
        cell.number_format = MONEY_FORMAT
        cell.alignment = _RIGHT
        if emphasis:
            label_cell.font = _BOLD
            cell.font = _BOLD
        self.ws.cell(row=self.row, column=3, value=key)
        self._name(key, self.row)
        return self.row

    def note(self, text: str, *, end_column: int = 2) -> None:
        self.row += 1
        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.font = _MUTED
        cell.alignment = _WRAP
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=end_column)

    def blank(self) -> None:
        self.row += 1

    def ref(self, key: str) -> str:
        return f"{self.INPUT}{self.cells[key]}"

    # --- the grid rows: the debt schedule, and the PFS schedule blocks ------

    def colhead(self, row: Row) -> None:
        self.row += 1
        for cell in row.cells:
            target = self.ws.cell(row=self.row, column=cell.c, value=cell.label)
            target.font = _HEADING_FONT
            target.fill = _HEADING_FILL
            target.alignment = _WRAP
        if not self.keys_beside_inputs:
            # The schema key under its label, on a hidden row.
            for index, key in enumerate(DEBT_COLUMNS, start=1):
                self.ws.cell(row=self.row + 1, column=index, value=key)
            self.ws.row_dimensions[self.row + 1].hidden = True

    def data(self, row: Row) -> None:
        self.row += 1
        if not self.keys_beside_inputs:
            self.ws.cell(row=self.row, column=self.key_column, value=f"r{row.ordinal}")
        for cell in row.cells:
            target = self.ws.cell(row=self.row, column=cell.c)
            target.protection = _UNLOCKED
            target.fill = _INPUT_FILL
            if cell.type == "money":
                target.number_format = MONEY_FORMAT
            elif cell.type == "rate":
                target.number_format = RATE_FORMAT
            elif cell.type == "date":
                target.number_format = DATE_FORMAT
            elif cell.type == "select" and cell.options:
                self._validation(tuple(cell.options)).add(target)
            self._name(cell.xlsx_name, self.row, cell.c)

    def totals(self, row: Row) -> None:
        self.row += 1
        for cell in row.cells:
            if cell.type == "label":
                self.ws.cell(row=self.row, column=cell.c, value=cell.label).font = _BOLD
                continue
            target = self.ws.cell(row=self.row, column=cell.c, value=self.resolve(cell.formula))
            target.number_format = MONEY_FORMAT
            target.font = _BOLD
            self._name(cell.xlsx_name, self.row, cell.c)

    def finish(self) -> None:
        self.ws.print_area = f"A1:{get_column_letter(self.last_column)}{self.row}"
        if self.last_column > 2:
            self.ws.page_setup.orientation = "landscape"
        self.ws.protection.sheet = True
        if self.keys_beside_inputs:
            self.ws.sheet_view.showGridLines = True


def _render(wb: Workbook, sheet: Sheet, *, first: bool = True) -> None:
    """Write one layout's rows into a sheet of `wb`, in row order."""
    form = _Form(wb, sheet, first=first)
    for row in sheet.rows:
        form.row = row.r - 1
        if row.kind == "title":
            form.line(row.label, _TITLE_FONT)
        elif row.kind == "subtitle":
            form.line(row.label, _MUTED)
        elif row.kind == "blank":
            form.blank()
        elif row.kind == "heading":
            form.heading(row.label, form.band(row.cells[0].colspan))
        elif row.kind == "note":
            form.note(row.label, end_column=form.band(row.cells[0].colspan)[-1])
        elif row.kind == "field":
            cell = next(cell for cell in row.cells if cell.key is not None)
            if cell.type == "money":
                form.money_row(cell.key, row.label, hint=cell.hint)
            else:
                form.text_field(cell.key, row.label, input=cell.type, options=cell.options, hint=cell.hint)
        elif row.kind == "formula" and row.block is None:
            cell = next(cell for cell in row.cells if cell.type == "formula")
            form.formula_row(cell.xlsx_name, row.label, form.resolve(cell.formula), emphasis=cell.emphasis)
        elif row.kind == "formula":
            form.totals(row)
        elif row.kind == "colhead":
            form.colhead(row)
        elif row.kind == "data":
            form.data(row)
        else:
            raise ValueError(f"unknown row kind {row.kind!r}")
    form.finish()


def _build(wb: Workbook, kind: str, *, first: bool = True) -> None:
    _render(wb, sheet_layout.layout(kind), first=first)


#: `dcterms:modified` inside docProps/core.xml, whatever openpyxl put there.
_MODIFIED_STAMP = re.compile(
    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)"
)

_PINNED_ISO = b"1980-01-01T00:00:00Z"


def _pin_zip_timestamps(raw: bytes) -> bytes:
    """openpyxl stamps zip entries with the wall clock. Rewrite them with one
    fixed timestamp so the same workbook is the same bytes on every build.

    docProps/core.xml needs the same treatment for a reason that is easy to
    miss: setting `wb.properties.modified` before saving does nothing, because
    openpyxl's writer overwrites that field with `datetime.now()` on its way
    out (writer/excel.py). Pinning it here — after the save, on the way through
    the zip — is the only place it sticks. Without this the bytes changed on
    every build, which made two determinism tests pass on luck alone: they
    compared two builds that usually, but not always, landed in the same
    wall-clock second.
    """
    source = zipfile.ZipFile(BytesIO(raw))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            pinned = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            pinned.compress_type = zipfile.ZIP_DEFLATED
            body = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = _MODIFIED_STAMP.sub(rb"\g<1>" + _PINNED_ISO + rb"\g<2>", body)
            target.writestr(pinned, body)
    return out.getvalue()


def _blank_workbook(title: str) -> Workbook:
    """An empty workbook with the pinned properties every download carries."""
    wb = Workbook()
    wb.properties.creator = _AUTHOR
    wb.properties.lastModifiedBy = _AUTHOR
    wb.properties.title = title
    wb.properties.created = _PINNED_STAMP
    wb.properties.modified = _PINNED_STAMP
    return wb


def _finished_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return _pin_zip_timestamps(buffer.getvalue())


@functools.cache
def build_workbook(kind: str) -> bytes:
    """The workbook for one kind, as bytes. Built once per process."""
    if kind not in sheet_layout.KINDS:
        raise KeyError(kind)
    wb = _blank_workbook(SHEET_TITLES[kind])
    _build(wb, kind)
    return _finished_bytes(wb)


@functools.cache
def build_packet_workbook() -> bytes:
    """All four forms as four tabs of one workbook, as bytes.

    The "one sheet we can forward the client or their accountant": a client or
    their accountant fills one file instead of four. Same renderer, same rows,
    same defined names — the prefixes (`pl.` / `bs.` / `ds.` / `pfs.`) and the
    sheet title inside each name keep them apart in one workbook. Four sheets
    of 50/56/22/80 rows sits inside the analyzer's budget, so a filled copy
    uploaded back is still read whole. Built once per process.
    """
    wb = _blank_workbook(PACKET_TITLE)
    for index, kind in enumerate(PACKET_KINDS):
        _build(wb, kind, first=index == 0)
    return _finished_bytes(wb)


def workbook_for_slug(slug: str) -> tuple[str, bytes] | None:
    """(attachment filename, bytes) for a public slug, or None for an unknown one."""
    if slug == PACKET_SLUG:
        return PACKET_ATTACHMENT_FILENAME, build_packet_workbook()
    kind = SLUGS.get(slug)
    if kind is None:
        return None
    return ATTACHMENT_FILENAMES[kind], build_workbook(kind)


def schema_keys(kind: str) -> list[str]:
    """Every key the workbook of this kind must carry a defined name for."""
    if kind in bss.KINDS:
        schema = bss.SCHEMA_FOR[kind]
        keys = [field.key for field in schema.header]
        for section in schema.sections:
            keys.extend(row.key for row in section.rows)
            keys.append(section.subtotal_key)
        return keys
    if kind == "pfs":
        keys = ["name", "business_name", "home_address", "business_phone", "statement_date"]
        for rows in (
            pfs_schema.ASSET_ROWS,
            pfs_schema.LIABILITY_ROWS,
            pfs_schema.INCOME_ROWS,
            pfs_schema.CONTINGENT_ROWS,
        ):
            keys.extend(row.key for row in rows)
        keys.extend(["total_assets", "total_liabilities", "net_worth", "total_income", "total_contingent"])
        return keys
    if kind == "debt_schedule":
        return ["business_name"] + [
            f"r{number}.{column}" for number in range(1, DEBT_ROWS + 1) for column in DEBT_COLUMNS
        ]
    raise KeyError(kind)
